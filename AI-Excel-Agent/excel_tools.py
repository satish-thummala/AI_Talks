
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

PROJECT_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = PROJECT_DIR / "AI_Talks_Sales_Workbook.xlsx"
SALES_SHEET = "SalesData"
REPORT_SHEET = "AIReport"


def _load_sales_data() -> pd.DataFrame:
    """Load and validate the source sales worksheet."""
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Workbook not found: {WORKBOOK_PATH}. Place "
            "AI_Talks_Sales_Workbook.xlsx beside excel_tools.py."
        )

    df = pd.read_excel(WORKBOOK_PATH, sheet_name=SALES_SHEET)

    required = {
        "OrderID", "OrderDate", "Product", "Category", "Region",
        "SalesRep", "Quantity", "UnitPrice", "Revenue", "OrderStatus",
    }
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"SalesData is missing columns: {sorted(missing)}")

    for column in ("Quantity", "UnitPrice", "Revenue"):
        df[column] = pd.to_numeric(df[column], errors="coerce")

    if df[["Quantity", "UnitPrice", "Revenue"]].isna().any().any():
        raise ValueError("Quantity, UnitPrice, and Revenue must be numeric.")

    return df


def get_sales_overview() -> dict[str, Any]:
    """Return headline metrics for the workbook."""
    df = _load_sales_data()
    return {
        "order_count": int(df["OrderID"].nunique()),
        "total_units_sold": int(df["Quantity"].sum()),
        "total_revenue": round(float(df["Revenue"].sum()), 2),
        "average_order_revenue": round(
            float(df.groupby("OrderID")["Revenue"].sum().mean()), 2
        ),
        "order_status_counts": {
            str(key): int(value)
            for key, value in df.groupby("OrderStatus")["OrderID"].nunique().items()
        },
    }


def get_product_performance(
    top_n: int = 3, lowest: bool = False
) -> list[dict[str, Any]]:
    """Rank products by revenue; lowest=True returns lowest-revenue products."""
    if not isinstance(top_n, int) or not 1 <= top_n <= 20:
        raise ValueError("top_n must be an integer between 1 and 20.")

    df = _load_sales_data()
    result = (
        df.groupby("Product", as_index=False)
        .agg(
            Revenue=("Revenue", "sum"),
            UnitsSold=("Quantity", "sum"),
            OrderCount=("OrderID", "nunique"),
        )
        .sort_values("Revenue", ascending=lowest)
        .head(top_n)
    )

    return [
        {
            "product": str(row.Product),
            "revenue": round(float(row.Revenue), 2),
            "units_sold": int(row.UnitsSold),
            "order_count": int(row.OrderCount),
        }
        for row in result.itertuples(index=False)
    ]


def get_region_performance() -> list[dict[str, Any]]:
    """Compare revenue, units, and order counts across regions."""
    df = _load_sales_data()
    total_revenue = float(df["Revenue"].sum())

    result = (
        df.groupby("Region", as_index=False)
        .agg(
            Revenue=("Revenue", "sum"),
            UnitsSold=("Quantity", "sum"),
            OrderCount=("OrderID", "nunique"),
        )
        .sort_values("Revenue", ascending=False)
    )

    return [
        {
            "region": str(row.Region),
            "revenue": round(float(row.Revenue), 2),
            "revenue_share_percent": round(
                float(row.Revenue) / total_revenue * 100, 2
            ) if total_revenue else 0.0,
            "units_sold": int(row.UnitsSold),
            "order_count": int(row.OrderCount),
        }
        for row in result.itertuples(index=False)
    ]


def get_order_status_summary() -> list[dict[str, Any]]:
    """Summarize order counts and revenue by order status."""
    df = _load_sales_data()

    result = (
        df.groupby("OrderStatus", as_index=False)
        .agg(
            OrderCount=("OrderID", "nunique"),
            Revenue=("Revenue", "sum"),
        )
        .sort_values("OrderStatus")
    )

    return [
        {
            "status": str(row.OrderStatus),
            "order_count": int(row.OrderCount),
            "revenue": round(float(row.Revenue), 2),
        }
        for row in result.itertuples(index=False)
    ]


def create_ai_report(
    summary: str, recommendations: list[str]
) -> dict[str, str]:
    """Create or replace AIReport without changing the SalesData worksheet."""
    if not isinstance(summary, str) or not summary.strip():
        raise ValueError("summary must be a non-empty string.")
    if not isinstance(recommendations, list):
        raise ValueError("recommendations must be a list of strings.")

    wb = load_workbook(WORKBOOK_PATH)

    if REPORT_SHEET in wb.sheetnames:
        del wb[REPORT_SHEET]

    ws = wb.create_sheet(REPORT_SHEET)

    ws["A1"] = "AI Excel Agent — Sales Analysis Report"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="174A7E")
    ws.merge_cells("A1:B1")

    ws["A3"] = "Generated summary"
    ws["A3"].font = Font(bold=True, color="FFFFFF")
    ws["A3"].fill = PatternFill("solid", fgColor="168A8A")

    ws["A4"] = summary.strip()
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:B4")
    ws.row_dimensions[4].height = max(
        45, min(180, 18 * (summary.count("\n") + 2))
    )

    ws["A6"] = "Recommendations"
    ws["A6"].font = Font(bold=True, color="FFFFFF")
    ws["A6"].fill = PatternFill("solid", fgColor="168A8A")

    for idx, recommendation in enumerate(recommendations, start=7):
        ws.cell(idx, 1, f"{idx - 6}.")
        ws.cell(idx, 2, str(recommendation))
        ws.cell(idx, 2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        ws.row_dimensions[idx].height = 32

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90
    ws.freeze_panes = "A3"

    wb.save(WORKBOOK_PATH)

    return {
        "status": "success",
        "message": f"Created the {REPORT_SHEET} worksheet.",
        "workbook": str(WORKBOOK_PATH),
    }


if __name__ == "__main__":
    # Test the tools without making a Groq API call.
    print("Sales overview:")
    print(get_sales_overview())

    print("\nThree lowest-revenue products:")
    print(get_product_performance(top_n=3, lowest=True))

    print("\nRegional performance:")
    print(get_region_performance())

    print("\nOrder status summary:")
    print(get_order_status_summary())
